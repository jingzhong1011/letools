# Regenerates inst/extdata/taiwan_lt_2023.csv from the Ministry of the
# Interior's 2023 (民國112年) national life table workbook, sheet 表1.
#
# Source: Department of Statistics, Ministry of the Interior, Taiwan.
# The workbook is not committed. Download it, then run from anywhere:
#
#     python3 data-raw/build_taiwan_lt.py [path/to/112年全國web.xlsx]
#
# With no argument, the workbook is looked for next to this script.
#
# The published table has qx, lx, dx, Lx, Tx and ex but no mx or ax. Both are
# recovered from identities (mx = dx / Lx, Lx = l(x+1) + ax * dx) so that
# lt(mx, ax = ax) can reproduce the published table exactly.

from pathlib import Path
import csv
import sys

import numpy as np
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
DEFAULT_XLSX = HERE / '112年全國web.xlsx'
OUT_CSV = HERE.parent / 'inst' / 'extdata' / 'taiwan_lt_2023.csv'

xlsx = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_XLSX
if not xlsx.is_file():
    sys.exit(f"Workbook not found: {xlsx}\n"
             "Download it from the Ministry of the Interior and pass its path.")

wb = load_workbook(xlsx)
rows = list(wb['表1'].iter_rows(values_only=True))

SEX = {'全體': 'total', '男性': 'male', '女性': 'female'}
current, data = None, {v: [] for v in SEX.values()}

for r in rows:
    a = r[0]
    if a is None: continue
    s = str(a).strip()
    if s in SEX:
        current = SEX[s]; continue
    if current is None: continue
    # keep integer ages and the open group only; drop the 0M..6M sub-year rows
    if not (s.isdigit() or s == '85+'): continue
    age = 85 if s == '85+' else int(s)
    qx, lx, dx, Lx, Tx, ex = (float(x) for x in r[1:7])
    data[current].append((age, qx, lx, dx, Lx, Tx, ex))

out = []
for sex, recs in data.items():
    recs.sort(key=lambda t: t[0])
    ages = [t[0] for t in recs]
    assert ages == list(range(86)), (sex, len(ages))
    arr = np.array([t[1:] for t in recs])
    qx, lx, dx, Lx, Tx, ex = arr.T
    n = len(ages)

    # ax from the person-years identity: Lx = 1*l(x+1) + ax*dx  (n = 1)
    ax = np.empty(n)
    ax[:-1] = (Lx[:-1] - lx[1:]) / dx[:-1]
    ax[-1] = Lx[-1] / lx[-1]          # open group: ax = e_open
    mx = dx / Lx

    for i in range(n):
        out.append(dict(sex=sex, age=ages[i], mx=mx[i], ax=ax[i], qx=qx[i],
                        lx=lx[i], dx=dx[i], Lx=Lx[i], Tx=Tx[i], ex=ex[i]))

with open(OUT_CSV, 'w', newline='') as f:
    w = csv.DictWriter(f, fieldnames=list(out[0].keys()))
    w.writeheader()
    for row in out:
        w.writerow({k: (v if isinstance(v, (str, int)) else format(float(v), '.15g'))
                    for k, v in row.items()})

print(f"{len(out)} rows written to {OUT_CSV}")
for sex in data:
    e0 = [r['ex'] for r in out if r['sex'] == sex][0]
    a0 = [r['ax'] for r in out if r['sex'] == sex][0]
    print(f"  {sex:7s} e0 = {e0:.4f}   a0 = {a0:.4f}")
